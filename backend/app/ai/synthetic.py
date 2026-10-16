"""Small made-up inputs (notes, Word spec, PNG label) used to record and test extraction of the newer file types."""

from pathlib import Path

import docx
import pymupdf

NOTES = """Product: Oat Crunch Granola
Brand: Meadow Mill
Net weight: 340 g
Allergens: contains oats and almonds
Store in a cool dry place
"""


def make_notes(path: Path) -> Path:
    path.write_text(NOTES, encoding="utf-8")
    return path


def make_docx(path: Path) -> Path:
    d = docx.Document()
    d.add_heading("Product specification", level=1)
    d.add_paragraph("Product name: Oat Crunch Granola")
    d.add_paragraph("Ingredients: oats, honey, almonds, sunflower oil, salt")
    d.add_heading("Nutrition per serving", level=2)
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Nutrient", "Amount"), ("Sodium", "140 mg"), ("Protein", "4 g")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_paragraph("Net weight: 340 g")
    d.save(str(path))
    return path


def make_png(path: Path) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=200)
    for y, text, size in ((50, "MEADOW MILL", 30), (100, "Oat Crunch Granola", 24), (150, "Net weight 340 g", 20)):
        page.insert_text((30, y), text, fontsize=size, fontname="helv")
    page.get_pixmap(dpi=120).save(str(path))
    doc.close()
    return path
