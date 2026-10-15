"""Word, plain-PDF, notes and image readers on tiny files built here."""

import docx
import pymupdf

from app.readers import read_file, read_text
from app.readers.pdf_document import is_dieline

def _png(w=3, h=2) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.clear_with(200)
    return pix.tobytes("png")


def test_docx_paragraphs_headings_tables(tmp_path):
    d = docx.Document()
    d.add_heading("Ingredients", level=1)
    d.add_paragraph("Oats, honey, salt")
    d.add_paragraph("   ")
    d.add_heading("Claims", level=2)
    d.add_paragraph("Good source of fiber")
    t = d.add_table(rows=2, cols=2)
    for (r, c), v in {(0, 0): "Nutrient", (0, 1): "Amount", (1, 0): "Fat", (1, 1): "3 g"}.items():
        t.cell(r, c).text = v
    p = tmp_path / "spec.docx"
    d.save(p)
    sf, pieces = read_file(p, "w")
    assert sf.type == "docx" and sf.status == "read"
    by = {x.text: x for x in pieces}
    assert by["Oats, honey, salt"].location.kind == "doc"
    assert by["Oats, honey, salt"].location.paragraph == 2
    assert by["Oats, honey, salt"].label == "Ingredients"
    assert by["Good source of fiber"].location.paragraph == 4  # empty paragraph skipped, headings count
    assert by["Good source of fiber"].label == "Claims"
    loc = by["3 g"].location
    assert (loc.table, loc.row, loc.col) == (1, 2, 2)
    assert by["3 g"].label == "Amount"


def _pdf(tmp_path, layers=()):
    doc = pymupdf.open()
    page = doc.new_page(width=300, height=300)
    page.insert_text((30, 50), "Net weight 340 g", fontsize=12)
    page.insert_text((30, 80), "Store in a cool dry place", fontsize=10)
    for name in layers:
        doc.add_ocg(name)
    p = tmp_path / "x.pdf"
    doc.save(p)
    return p


def test_plain_pdf_is_document(tmp_path):
    p = _pdf(tmp_path)
    assert not is_dieline(p)
    sf, pieces = read_file(p, "d")
    assert sf.type == "pdf" and sf.pages == 1
    assert [x.text for x in pieces] == ["Net weight 340 g", "Store in a cool dry place"]
    loc = pieces[0].location
    assert loc.kind == "pdf" and loc.page == 1 and loc.panel is None and loc.bbox_mm is not None
    assert all(x.status == "visible" for x in pieces)


def test_pdf_with_artwork_panel_layer_is_dieline(tmp_path):
    assert is_dieline(_pdf(tmp_path, layers=("Artwork panel", "DESIGN")))


def test_notes_line_numbers(tmp_path):
    p = tmp_path / "n.md"
    p.write_text("# Notes\n\nNo artificial colours\n  \nLaunch in India\n", encoding="utf-8")
    sf, pieces = read_file(p, "n")
    assert sf.type == "text"
    assert [(x.text, x.location.line) for x in pieces] == [("# Notes", 1), ("No artificial colours", 3), ("Launch in India", 5)]
    assert all(x.location.kind == "text" for x in pieces)


def test_pasted_text():
    sf, pieces = read_text("a\n\nb", "pp", "Pasted notes")
    assert sf.name == "Pasted notes" and sf.type == "text"
    assert [x.location.line for x in pieces] == [1, 3]


def test_image_single_vision_piece(tmp_path):
    p = tmp_path / "pack.png"
    p.write_bytes(_png())
    sf, pieces = read_file(p, "i")
    assert (sf.type, sf.width_px, sf.height_px, sf.status) == ("image", 3, 2, "read")
    assert len(pieces) == 1
    x = pieces[0]
    assert x.text == "" and x.status == "visible" and x.reason == "read by vision"
    assert x.location.kind == "image" and x.location.bbox_mm[:2] == (0.0, 0.0)
