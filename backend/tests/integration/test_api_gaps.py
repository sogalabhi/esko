"""API gaps the UI needs: raw file route, file types and errors, skip / undo / resolve (server state survives a re-GET)."""

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app import store
from app.main import app
from app.routers import files, project as project_router
from tests.integration.test_project_facts import CARTON, SHEET, fake_client


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "UPLOAD_DIR", tmp_path / "uploads")
    store.clear()
    app.dependency_overrides[project_router.get_ai_client] = fake_client
    yield TestClient(app)
    app.dependency_overrides.clear()
    store.clear()


@pytest.fixture
def run(client):
    for p in (SHEET, CARTON):
        client.post("/api/files", files={"file": (p.name, p.read_bytes())})
    return client.post("/api/project/run").json()


def upload(client, name, data):
    r = client.post("/api/files", files={"file": (name, data)})
    assert r.status_code == 200, r.text
    return r.json()


def file_of(client, file_id):
    return next(f for f in client.get("/api/project").json()["files"] if f["id"] == file_id)


# --- raw ------------------------------------------------------------------------------------------


def test_raw_streams_the_stored_bytes_with_content_type(client):
    up = upload(client, CARTON.name, CARTON.read_bytes())
    r = client.get(f"/api/files/{up['id']}/raw")
    assert r.status_code == 200 and r.content == CARTON.read_bytes()
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"].startswith("inline")
    sheet = upload(client, SHEET.name, SHEET.read_bytes())
    assert "spreadsheetml" in client.get(f"/api/files/{sheet['id']}/raw").headers["content-type"]


def test_raw_head_has_headers_and_no_body(client):
    up = upload(client, CARTON.name, CARTON.read_bytes())
    r = client.head(f"/api/files/{up['id']}/raw")
    assert r.status_code == 200 and r.content == b""
    assert r.headers["content-type"] == "application/pdf"
    assert int(r.headers["content-length"]) == len(CARTON.read_bytes())


def test_raw_unknown_id_is_404_with_a_real_reason(client):
    for call in (client.get, client.head):
        assert call("/api/files/nope/raw").status_code == 404
    assert client.get("/api/files/nope/raw").json()["detail"] == "file not found"


# --- file errors and types --------------------------------------------------------------------------


def test_file_error_is_passed_through(client):
    bad = upload(client, "a.pdf", b"%PDF-1.4 fake")
    f = file_of(client, bad["id"])
    assert f["status"] == "error" and f["error"] and f["error"] == bad["error"]
    ok = upload(client, SHEET.name, SHEET.read_bytes())
    assert "error" not in file_of(client, ok["id"])


def test_file_types_and_counts_for_each_reader(client, tmp_path):
    carton = upload(client, CARTON.name, CARTON.read_bytes())
    sheet = upload(client, SHEET.name, SHEET.read_bytes())

    d = docx.Document()
    d.add_heading("Ingredients", level=1)
    d.add_paragraph("Oats, honey, salt")
    d.add_paragraph("Good source of fiber")
    d.save(tmp_path / "spec.docx")
    word = upload(client, "spec.docx", (tmp_path / "spec.docx").read_bytes())

    note = upload(client, "notes.txt", b"Net weight 340 g\n\nMade in USA\nKeep dry\n")

    pdf = pymupdf.open()
    pdf.new_page().insert_text((30, 50), "Net weight 340 g", fontsize=12)
    pdf.new_page().insert_text((30, 50), "Store dry", fontsize=12)
    plain = upload(client, "spec.pdf", pdf.tobytes())

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 3, 2), False)
    pix.clear_with(200)
    image = upload(client, "label.png", pix.tobytes("png"))

    types = {f["id"]: (f["type"], f["pages"]) for f in client.get("/api/project").json()["files"]}
    assert types[carton["id"]] == ("pdf", carton["pages"])
    sheets = next(f for f in client.get("/api/files").json() if f["id"] == sheet["id"])["sheets"]
    assert types[sheet["id"]] == ("sheet", len(sheets)) and len(sheets) > 0
    assert types[word["id"]] == ("docx", 3)
    assert types[note["id"]][0] == "text" and types[note["id"]][1] >= 3
    assert types[plain["id"]] == ("document", 2)
    assert types[image["id"]] == ("image", 1)


def test_run_keeps_the_real_types(client, run):
    assert {f["type"] for f in run["files"]} == {"pdf", "sheet"}
    assert run["files"] == client.get("/api/project").json()["files"]


# --- skip / undo / resolve ----------------------------------------------------------------------------


def q_of(project, qid):
    return next(q for q in project["questions"] if q["id"] == qid)


def test_skip_then_undo_survive_a_reget(client, run):
    qid = next(q["id"] for q in run["questions"] if q["status"] == "open")
    r = client.post(f"/api/questions/{qid}/skip")
    assert r.status_code == 200
    assert q_of(r.json(), qid)["status"] == "skipped"
    assert q_of(client.get("/api/project").json(), qid)["status"] == "skipped"
    assert client.post(f"/api/questions/{qid}/skip").status_code == 409

    u = client.post(f"/api/questions/{qid}/undo")
    assert u.status_code == 200 and q_of(u.json(), qid)["status"] == "open"
    assert q_of(client.get("/api/project").json(), qid)["status"] == "open"
    assert client.post(f"/api/questions/{qid}/undo").status_code == 409  # nothing left to undo
    assert client.post("/api/questions/q-none/skip").status_code == 404


def test_undo_an_answer_restores_the_project_as_it_was(client, run):
    before = client.get("/api/project").json()
    vegan = next(q for q in run["questions"] if "vegan" in q["text"].lower())
    label = next(a for a in vegan["answers"] if "honey" in a.lower() and "keep" in a.lower())
    answered = client.post("/api/answers", json={"questionId": vegan["id"], "answer": label}).json()
    assert answered != before and any(c.get("resolution") for c in answered["conflicts"])

    undone = client.post(f"/api/questions/{vegan['id']}/undo").json()
    assert undone == before  # every effect of the answer is gone
    assert client.get("/api/project").json() == before


def test_undo_keeps_other_decisions(client, run):
    first, second = [q for q in run["questions"] if q["status"] == "open"][:2]
    client.post(f"/api/questions/{first['id']}/skip")
    client.post(f"/api/questions/{second['id']}/skip")
    got = client.post(f"/api/questions/{first['id']}/undo").json()
    assert q_of(got, first["id"])["status"] == "open" and q_of(got, second["id"])["status"] == "skipped"


def test_resolve_a_conflict_is_logged_as_the_person_who_decided(client, run):
    cid = next(c["id"] for c in run["conflicts"] if "resolution" not in c)
    r = client.post(f"/api/conflicts/{cid}/resolve", json={"choice": "left", "by": "Priya"})
    assert r.status_code == 200, r.text
    res = next(c for c in r.json()["conflicts"] if c["id"] == cid)["resolution"]
    assert res["choice"] == "left" and res["by"] == "Priya" and res["at"]
    again = next(c for c in client.get("/api/project").json()["conflicts"] if c["id"] == cid)
    assert again["resolution"] == res
    assert client.post(f"/api/conflicts/{cid}/resolve", json={"choice": "right", "by": "x"}).status_code == 409


def test_resolve_custom_needs_text_and_keeps_it(client, run):
    cid = next(c["id"] for c in run["conflicts"] if "resolution" not in c)
    assert client.post(f"/api/conflicts/{cid}/resolve", json={"choice": "custom", "by": "Priya"}).status_code == 422
    assert client.post(f"/api/conflicts/{cid}/resolve", json={"choice": "bogus", "by": "Priya"}).status_code == 422
    assert client.post(f"/api/conflicts/{cid}/resolve", json={"choice": "intended", "by": " "}).status_code == 422
    r = client.post(f"/api/conflicts/{cid}/resolve", json={"choice": "custom", "text": "Use 340 g", "by": "Priya"})
    assert r.status_code == 200
    res = next(c for c in r.json()["conflicts"] if c["id"] == cid)["resolution"]
    assert (res["choice"], res["text"], res["by"]) == ("custom", "Use 340 g", "Priya")
    assert any(f["value"] == "Use 340 g" for f in client.get("/api/project").json()["facts"])
    assert client.post("/api/conflicts/c-none/resolve", json={"choice": "intended", "by": "x"}).status_code == 404


def test_decisions_without_a_run_are_404(client):
    assert client.post("/api/questions/q/skip").status_code == 404
    assert client.post("/api/conflicts/c/resolve", json={"choice": "intended", "by": "x"}).status_code == 404
    assert client.post("/api/questions/q/undo").status_code == 404
